import datetime
from pathlib import Path
from typing import Optional
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from .models import DocumentStructure, DocumentSection, DocumentTable, DocumentCodeBlock


# Brand Colors
NAVY_PRIMARY = RGBColor(30, 58, 138)      # #1E3A8A
SLATE_ACCENT = RGBColor(37, 99, 235)      # #2563EB
TEXT_PRIMARY = RGBColor(31, 41, 55)       # #1F2937
TEXT_MUTED = RGBColor(100, 116, 139)      # #64748B
BG_LIGHT = "F8FAFC"
BG_CODE = "F1F5F9"
BORDER_COLOR = "CBD5E1"
NAVY_HEX = "1E3A8A"


def set_cell_background(cell, hex_color: str):
    """Sets background shading of a docx table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shd if present
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Sets internal cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, hex_color="E2E8F0"):
    """Applies subtle border styling to an entire table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{hex_color}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{hex_color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def generate_docx(doc_struct: DocumentStructure, output_path: str) -> str:
    """
    Generates a high-quality, professional Word document (.docx)
    from a structured DocumentStructure model.
    """
    doc = Document()

    # 1. Page Configuration: Standard Letter with 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

        # Header: Document Title & Sovereign tag
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header_run = header_p.add_run(f"SOVA Sovereign Document Intelligence • {doc_struct.title[:45]}")
        header_run.font.name = "Calibri"
        header_run.font.size = Pt(8.5)
        header_run.font.color.rgb = TEXT_MUTED

        # Footer: Confidentiality & Page Notice
        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        confidentiality = doc_struct.confidentiality or "Confidential - Air-Gapped On-Premise"
        footer_run = footer_p.add_run(f"{confidentiality} • Generated Locally by SOVA")
        footer_run.font.name = "Calibri"
        footer_run.font.size = Pt(8.5)
        footer_run.font.color.rgb = TEXT_MUTED

    # 2. Configure Default Styles
    style_normal = doc.styles["Normal"]
    style_normal.font.name = "Calibri"
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = TEXT_PRIMARY
    style_normal.paragraph_format.line_spacing = 1.2
    style_normal.paragraph_format.space_after = Pt(6)

    # 3. Document Title Block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(4)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run(doc_struct.title)
    title_run.font.name = "Segoe UI"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = NAVY_PRIMARY

    # Subtitle if present
    if doc_struct.subtitle:
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_before = Pt(0)
        sub_p.paragraph_format.space_after = Pt(12)
        sub_run = sub_p.add_run(doc_struct.subtitle)
        sub_run.font.name = "Calibri"
        sub_run.font.size = Pt(13)
        sub_run.font.italic = True
        sub_run.font.color.rgb = TEXT_MUTED

    # Metadata Strip (Author, Date, Organization)
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(14)
    author_text = doc_struct.author or "SOVA Sovereign Intelligence"
    date_text = doc_struct.date or datetime.date.today().strftime("%d %B %Y")
    org_text = doc_struct.organization or "SOVA Industrial Systems"
    
    meta_run1 = meta_p.add_run("Author: ")
    meta_run1.font.bold = True
    meta_run1.font.size = Pt(9.5)
    meta_run1.font.color.rgb = TEXT_MUTED
    meta_p.add_run(f"{author_text}   |   ")
    
    meta_run2 = meta_p.add_run("Date: ")
    meta_run2.font.bold = True
    meta_run2.font.size = Pt(9.5)
    meta_run2.font.color.rgb = TEXT_MUTED
    meta_p.add_run(f"{date_text}   |   ")

    meta_run3 = meta_p.add_run("Classification: ")
    meta_run3.font.bold = True
    meta_run3.font.size = Pt(9.5)
    meta_run3.font.color.rgb = TEXT_MUTED
    meta_p.add_run(f"{doc_struct.confidentiality or 'Air-Gapped'}")

    # Thematic Divider
    divider_p = doc.add_paragraph()
    divider_p.paragraph_format.space_after = Pt(14)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="{NAVY_HEX}"/></w:pBdr>')
    divider_p._p.get_or_add_pPr().append(pBdr)

    # 4. Executive Summary Callout Box (if present)
    if doc_struct.summary:
        sum_table = doc.add_table(rows=1, cols=1)
        sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sum_cell = sum_table.cell(0, 0)
        set_cell_background(sum_cell, "EFF6FF")  # Subtle light blue
        set_cell_margins(sum_cell, top=140, bottom=140, left=200, right=200)

        # Blue left accent border
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/>'
            f'<w:left w:val="single" w:sz="24" w:color="2563EB"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        sum_cell._tc.get_or_add_tcPr().append(tcBorders)

        sum_p = sum_cell.paragraphs[0]
        sum_p.paragraph_format.space_after = Pt(2)
        sum_title = sum_p.add_run("EXECUTIVE SUMMARY\n")
        sum_title.font.name = "Segoe UI"
        sum_title.font.bold = True
        sum_title.font.size = Pt(10)
        sum_title.font.color.rgb = SLATE_ACCENT

        sum_body = sum_p.add_run(doc_struct.summary)
        sum_body.font.name = "Calibri"
        sum_body.font.size = Pt(10)
        sum_body.font.color.rgb = TEXT_PRIMARY

        # Space after summary
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(8)

    # 5. Process Sections
    for section_idx, sec in enumerate(doc_struct.sections):
        # Section Heading
        h_p = doc.add_paragraph()
        h_p.paragraph_format.keep_with_next = True
        h_run = h_p.add_run(sec.heading)
        h_run.font.name = "Segoe UI"

        if sec.level == 1:
            h_p.paragraph_format.space_before = Pt(16 if section_idx > 0 else 8)
            h_p.paragraph_format.space_after = Pt(6)
            h_run.font.size = Pt(15)
            h_run.font.bold = True
            h_run.font.color.rgb = NAVY_PRIMARY
        elif sec.level == 2:
            h_p.paragraph_format.space_before = Pt(12)
            h_p.paragraph_format.space_after = Pt(4)
            h_run.font.size = Pt(12.5)
            h_run.font.bold = True
            h_run.font.color.rgb = SLATE_ACCENT
        else:
            h_p.paragraph_format.space_before = Pt(10)
            h_p.paragraph_format.space_after = Pt(3)
            h_run.font.size = Pt(11)
            h_run.font.bold = True
            h_run.font.color.rgb = TEXT_PRIMARY

        # Paragraphs
        for para_text in sec.paragraphs:
            if not para_text.strip():
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.2
            p_run = p.add_run(para_text.strip())
            p_run.font.name = "Calibri"
            p_run.font.size = Pt(10.5)

        # Bullet Points
        for bullet in sec.bullet_points:
            if not bullet.strip():
                continue
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_after = Pt(3)
            bp_run = bp.add_run(bullet.strip())
            bp_run.font.name = "Calibri"
            bp_run.font.size = Pt(10)

        # Numbered List
        for num_item in sec.numbered_list:
            if not num_item.strip():
                continue
            np = doc.add_paragraph(style="List Number")
            np.paragraph_format.space_after = Pt(3)
            np_run = np.add_run(num_item.strip())
            np_run.font.name = "Calibri"
            np_run.font.size = Pt(10)

        # Tables
        for tbl_data in sec.tables:
            if not tbl_data.headers and not tbl_data.rows:
                continue

            if tbl_data.caption:
                cap_p = doc.add_paragraph()
                cap_p.paragraph_format.space_before = Pt(8)
                cap_p.paragraph_format.space_after = Pt(3)
                cap_p.paragraph_format.keep_with_next = True
                cap_run = cap_p.add_run(f"Table: {tbl_data.caption}")
                cap_run.font.name = "Calibri"
                cap_run.font.size = Pt(9.5)
                cap_run.font.bold = True
                cap_run.font.color.rgb = SLATE_ACCENT

            num_cols = len(tbl_data.headers) if tbl_data.headers else len(tbl_data.rows[0])
            table = doc.add_table(rows=0, cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(table, hex_color=BORDER_COLOR)

            # Header Row
            if tbl_data.headers:
                hdr_row = table.add_row()
                # Repeat header on new pages
                trPr = hdr_row._tr.get_or_add_trPr()
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

                for c_idx, h_text in enumerate(tbl_data.headers):
                    cell = hdr_row.cells[c_idx]
                    set_cell_background(cell, NAVY_HEX)
                    set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    cp = cell.paragraphs[0]
                    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    crun = cp.add_run(str(h_text))
                    crun.font.name = "Segoe UI"
                    crun.font.bold = True
                    crun.font.size = Pt(9.5)
                    crun.font.color.rgb = RGBColor(255, 255, 255)

            # Data Rows
            for r_idx, row_values in enumerate(tbl_data.rows):
                data_row = table.add_row()
                trPr = data_row._tr.get_or_add_trPr()
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                is_even = (r_idx % 2 == 0)
                bg_color = BG_LIGHT if is_even else "FFFFFF"

                for c_idx in range(num_cols):
                    cell = data_row.cells[c_idx]
                    val = str(row_values[c_idx]) if c_idx < len(row_values) else ""
                    set_cell_background(cell, bg_color)
                    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_after = Pt(0)
                    crun = cp.add_run(val)
                    crun.font.name = "Calibri"
                    crun.font.size = Pt(9)
                    crun.font.color.rgb = TEXT_PRIMARY

            # Spacer after table
            post_tbl_p = doc.add_paragraph()
            post_tbl_p.paragraph_format.space_before = Pt(6)
            post_tbl_p.paragraph_format.space_after = Pt(4)

        # Code Blocks
        for code_block in sec.code_blocks:
            if not code_block.code.strip():
                continue
            code_table = doc.add_table(rows=1, cols=1)
            code_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            code_cell = code_table.cell(0, 0)
            set_cell_background(code_cell, BG_CODE)
            set_cell_margins(code_cell, top=120, bottom=120, left=180, right=180)

            # Gray border
            cBrd = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="6" w:color="{BORDER_COLOR}"/>'
                f'<w:left w:val="single" w:sz="6" w:color="{BORDER_COLOR}"/>'
                f'<w:bottom w:val="single" w:sz="6" w:color="{BORDER_COLOR}"/>'
                f'<w:right w:val="single" w:sz="6" w:color="{BORDER_COLOR}"/>'
                f'</w:tcBorders>'
            )
            code_cell._tc.get_or_add_tcPr().append(cBrd)

            code_p = code_cell.paragraphs[0]
            code_p.paragraph_format.space_after = Pt(0)
            code_run = code_p.add_run(code_block.code.strip())
            code_run.font.name = "Consolas"
            code_run.font.size = Pt(9)
            code_run.font.color.rgb = RGBColor(30, 41, 59)

            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(6)

        # Page break if requested for section
        if sec.page_break_after and section_idx < len(doc_struct.sections) - 1:
            doc.add_page_break()

    # Save to disk
    out_file = Path(output_path).resolve()
    out_file.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_file))

    return str(out_file)
