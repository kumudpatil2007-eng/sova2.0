import os
import sys
import asyncio
from pathlib import Path
from fastapi.testclient import TestClient

BACKEND_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BACKEND_DIR))

from app.main import app
from app.config import settings
from app.llm.router import route_query, TaskCategory, detect_document_generation_intent, is_general_query
from app.services.document_processor import process_document, get_cached_document
from app.services.rag_service import chunk_document, retrieve_relevant_context
from app.services.llm_service import generate_response, check_ollama_status
import docx
import pymupdf


def banner(text: str):
    print("\n" + "=" * 70)
    print(f" {text}")
    print("=" * 70)


async def run_tests():
    client = TestClient(app)

    banner("OLLAMA HEALTH & MODEL CHECK")
    status = await check_ollama_status()
    print(f"Ollama Online: {status['online']}")
    print(f"Available models: {status['available_models']}")
    assert status['online'], "Ollama must be running"
    assert "llama3.2:3b" in status['available_models'], "llama3.2:3b must be installed"
    assert "qwen2.5vl:3b" in status['available_models'], "qwen2.5vl:3b must be installed"

    # -------------------------------------------------------------
    # TEST 1: "2 + 2" -> llama3.2:3b -> normal answer
    # -------------------------------------------------------------
    banner("TEST 1: '2 + 2' -> llama3.2:3b -> normal answer")
    q1 = "What is 2 + 2?"
    decision_1 = route_query(q1)
    print(f"Query: '{q1}'")
    print(f"Routing Category: {decision_1.category}")
    print(f"Routed Model: {decision_1.model_name}")
    print(f"Requires Vision: {decision_1.requires_vision}")
    assert decision_1.model_name == "llama3.2:3b", f"Expected llama3.2:3b, got {decision_1.model_name}"
    assert decision_1.requires_vision is False, "2 + 2 must NOT require vision"

    res_1 = await generate_response(q1)
    print(f"Model used: {res_1['model']}")
    print(f"Answer: {res_1['answer'].strip()}")
    assert "4" in res_1["answer"], f"Expected '4' in answer, got: {res_1['answer']}"
    assert "llama3.2:3b" in res_1["model"]
    print(">>> TEST 1 PASSED!\n")

    # -------------------------------------------------------------
    # TEST 2: Upload a text PDF and ask a question
    # -> document retrieval works -> answer generated
    # -------------------------------------------------------------
    banner("TEST 2: Text PDF upload + question -> retrieval works -> answer generated")
    sample_pdf = BACKEND_DIR / "sample_data" / "Heat_Exchanger_E204_Inspection_Report.pdf"
    assert sample_pdf.exists(), f"Sample PDF not found at {sample_pdf}"

    doc_2 = await process_document(str(sample_pdf), sample_pdf.name, "suite_test_text_pdf")
    chunk_document(doc_2)
    print(f"Document processed: {doc_2['filename']} ({doc_2['total_pages']} pages, {len(doc_2['full_text'])} chars)")
    print(f"requires_vision: {doc_2['requires_vision']}")
    print(f"is_scanned: {doc_2['is_scanned']}")
    print(f"images_base64 count: {len(doc_2['images_base64'])}")
    assert doc_2['requires_vision'] is False, "Text inspection report must NOT require vision"
    assert len(doc_2['images_base64']) == 0, "Text report should not produce image frames"

    q2 = "What is the measured minimum thickness and what is the retirement thickness for E-204?"
    decision_2 = route_query(q2, doc_data=doc_2, doc_extension=".pdf", doc_has_images=False)
    print(f"Query: '{q2}'")
    print(f"Routing Category: {decision_2.category}")
    print(f"Routed Model: {decision_2.model_name}")
    print(f"Requires Vision: {decision_2.requires_vision}")
    assert decision_2.model_name == "llama3.2:3b", f"Expected llama3.2:3b for text PDF, got {decision_2.model_name}"
    assert decision_2.requires_vision is False

    retrieval_2 = retrieve_relevant_context("suite_test_text_pdf", q2, top_k=3)
    assert len(retrieval_2["sources"]) > 0, "Document retrieval should find relevant sources"
    print(f"Retrieved {len(retrieval_2['sources'])} source chunks.")

    res_2 = await generate_response(
        user_message=q2,
        context=retrieval_2["context_text"],
        doc_data=doc_2
    )
    print(f"Model used: {res_2['model']}")
    print(f"Answer: {res_2['answer'][:250]}...")
    assert "llama3.2:3b" in res_2["model"]
    assert any(k in res_2["answer"].lower() for k in ["thickness", "mm", "6.2", "4.8", "retirement", "e-204", "e204"]), \
        "Answer should be grounded in the inspection report"
    print(">>> TEST 2 PASSED!\n")

    # -------------------------------------------------------------
    # TEST 3: Upload a scanned/image PDF
    # -> qwen2.5vl:3b handles visual content -> no multimodal error from llama3.2:3b
    # -------------------------------------------------------------
    banner("TEST 3: Scanned/image PDF -> qwen2.5vl:3b -> NO multimodal error from llama3.2:3b")
    # Locate a drawing / P&ID or create a synthetic scanned page PDF
    pid_files = list(Path(BACKEND_DIR / "uploads").glob("**/SOVA_Test_P_and_ID.pdf"))
    if not pid_files:
        pid_files = list(Path(BACKEND_DIR.parent / "uploads").glob("**/SOVA_Test_P_and_ID.pdf"))

    if pid_files:
        test_visual_pdf = pid_files[0]
    else:
        # Create a 1-page visual PDF
        test_visual_pdf = BACKEND_DIR / "uploads" / "temp_scanned_test.pdf"
        test_visual_pdf.parent.mkdir(parents=True, exist_ok=True)
        doc = pymupdf.open()
        page = doc.new_page()
        page.draw_rect(pymupdf.Rect(50, 50, 400, 300), color=(0, 0, 1), fill=(0.9, 0.9, 1))
        page.insert_text(pymupdf.Point(70, 100), "DIAGRAM: PUMP P-101 TO VESSEL V-201", fontsize=14)
        doc.save(str(test_visual_pdf))
        doc.close()

    doc_3 = await process_document(str(test_visual_pdf), test_visual_pdf.name, "suite_test_visual_pdf")
    print(f"Document processed: {doc_3['filename']}")
    print(f"requires_vision: {doc_3['requires_vision']}")
    print(f"images_base64 count: {len(doc_3['images_base64'])}")
    assert doc_3['requires_vision'] is True, "Visual document must require vision"
    assert len(doc_3['images_base64']) > 0, "Visual document must have rendered frames"

    q3 = "Analyze this diagram and identify the equipment shown."
    decision_3 = route_query(q3, doc_data=doc_3, doc_extension=".pdf", doc_has_images=True)
    print(f"Query: '{q3}'")
    print(f"Routing Category: {decision_3.category}")
    print(f"Routed Model: {decision_3.model_name}")
    print(f"Requires Vision: {decision_3.requires_vision}")
    assert decision_3.model_name == "qwen2.5vl:3b", f"Expected qwen2.5vl:3b for visual document, got {decision_3.model_name}"
    assert decision_3.requires_vision is True

    # Run inference with qwen2.5vl:3b
    res_3 = await generate_response(
        user_message=q3,
        context=doc_3.get("full_text", ""),
        images_base64=doc_3["images_base64"],
        doc_data=doc_3
    )
    print(f"Model used: {res_3['model']}")
    print(f"Answer: {res_3['answer'][:250]}...")
    assert "qwen2.5vl:3b" in res_3["model"]
    assert "multimodal data provided" not in res_3["answer"].lower(), "Multimodal error must NOT occur"
    print(">>> TEST 3 PASSED!\n")

    # -------------------------------------------------------------
    # TEST 4: Upload Heat_Exchanger_E204_Inspection_Report.pdf
    # -> ask "Generate a DOCX report based on this document."
    # -> document analyzed -> DOCX generated -> download works
    # -------------------------------------------------------------
    banner("TEST 4: Heat Exchanger -> 'Generate a DOCX report' -> DOCX generated -> download works")
    cached_e204 = get_cached_document("suite_test_text_pdf")
    if not cached_e204:
        cached_e204 = await process_document(str(sample_pdf), sample_pdf.name, "suite_test_text_pdf")

    q4 = "Generate a DOCX report based on this document."
    decision_4 = route_query(q4, doc_data=cached_e204, doc_extension=".pdf", doc_has_images=len(cached_e204['images_base64']) > 0)
    print(f"Query: '{q4}'")
    print(f"Routing Category: {decision_4.category}")
    print(f"Routed Model: {decision_4.model_name}")
    print(f"Requires Vision: {decision_4.requires_vision}")
    print(f"Is Doc Gen: {decision_4.is_document_generation}")
    print(f"Requested Formats: {decision_4.requested_formats}")
    assert decision_4.category == TaskCategory.DOCUMENT_GENERATION
    assert decision_4.is_document_generation is True
    assert decision_4.model_name == "llama3.2:3b", f"Text PDF report generation should use llama3.2:3b, got {decision_4.model_name}"
    assert decision_4.requires_vision is False, "llama3.2:3b must NEVER have requires_vision=True"

    # Context is the full document text for document authoring
    res_4 = await generate_response(
        user_message=q4,
        context=cached_e204.get("full_text", "")[:3500],
        doc_data=cached_e204
    )
    print(f"Model used: {res_4['model']}")
    print(f"Response text: {res_4['answer'][:180]}...")
    assert "Document generated successfully" in res_4["answer"], "Response must indicate successful document generation"

    gen_doc_4 = res_4.get("generated_document")
    assert gen_doc_4 is not None, "generated_document metadata must be returned"
    assert "docx" in gen_doc_4["files"], "DOCX file must be present in generated_document"

    docx_info = gen_doc_4["files"]["docx"]
    docx_path = docx_info["local_path"]
    download_url = docx_info["download_url"]
    print(f"Generated DOCX Path: {docx_path} ({docx_info['size_formatted']})")
    print(f"Download URL: {download_url}")
    assert os.path.exists(docx_path), f"DOCX file {docx_path} does not exist on disk"
    assert os.path.getsize(docx_path) > 1000, "DOCX file is too small"

    # Inspect generated docx with python-docx
    doc_obj = docx.Document(docx_path)
    paragraphs = [p.text for p in doc_obj.paragraphs if p.text.strip()]
    full_docx_text = "\n".join(paragraphs)
    print(f"DOCX contains {len(doc_obj.paragraphs)} paragraphs and {len(doc_obj.tables)} tables.")
    assert len(doc_obj.tables) >= 1 or len(paragraphs) >= 4, "DOCX should have structured content"
    assert any(k in full_docx_text.lower() for k in ["heat exchanger", "e-204", "e204", "inspection", "thickness"]), \
        "DOCX content must be grounded in the Heat Exchanger inspection report"

    # Test download endpoint via FastAPI TestClient
    dl_resp = client.get(download_url)
    assert dl_resp.status_code == 200, f"Download failed with status {dl_resp.status_code}"
    assert "wordprocessingml" in dl_resp.headers.get("content-type", "")
    assert len(dl_resp.content) == os.path.getsize(docx_path)
    print(f"Download HTTP 200 OK verified ({len(dl_resp.content)} bytes received).")
    print(">>> TEST 4 PASSED!\n")

    # -------------------------------------------------------------
    # TEST 5: Ask a normal question after document processing
    # -> llama3.2:3b still works
    # -------------------------------------------------------------
    banner("TEST 5: Normal question after document processing -> llama3.2:3b still works")
    q5 = "Explain binary search."
    # Simulate doc_data still present in session
    decision_5 = route_query(q5, doc_data=cached_e204, doc_extension=".pdf", doc_has_images=False)
    print(f"Query: '{q5}'")
    print(f"Routing Category: {decision_5.category}")
    print(f"Routed Model: {decision_5.model_name}")
    print(f"Requires Vision: {decision_5.requires_vision}")
    assert decision_5.model_name == "llama3.2:3b", f"Expected llama3.2:3b, got {decision_5.model_name}"
    assert decision_5.requires_vision is False

    res_5 = await generate_response(
        user_message=q5,
        context="",
        doc_data=cached_e204
    )
    print(f"Model used: {res_5['model']}")
    print(f"Answer: {res_5['answer'][:200]}...")
    assert "llama3.2:3b" in res_5["model"]
    assert any(k in res_5["answer"].lower() for k in ["binary search", "sorted", "divide", "algorithm", "logarithmic", "half"]), \
        "Answer should correctly explain binary search"
    assert "not available in the supplied documents" not in res_5["answer"].lower(), \
        "General question should not be rejected due to document presence"
    print(">>> TEST 5 PASSED!\n")

    # -------------------------------------------------------------
    # ERROR HANDLING TEST: When multimodal model is unavailable
    # -------------------------------------------------------------
    banner("ERROR HANDLING TEST: Unavailable multimodal model error message")
    from app.llm.local_client import OllamaModelNotFoundError
    from app.llm.local_client import ollama_client
    try:
        await ollama_client.resolve_available_model("qwen2.5vl:nonexistent_model_tag")
        assert False, "Should have raised OllamaModelNotFoundError"
    except OllamaModelNotFoundError as e:
        print(f"Caught expected error: {e}")
        assert "Multimodal local model unavailable. Please install qwen2.5vl:3b using Ollama." in str(e)
    print(">>> ERROR HANDLING TEST PASSED!\n")

    banner("ALL 5 REQUIRED TEST CASES + ERROR HANDLING TEST PASSED SUCCESSFULLY!")


if __name__ == "__main__":
    asyncio.run(run_tests())
